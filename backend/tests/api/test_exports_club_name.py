"""The member, licence and insurance exports name the club instead of showing its id."""

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from src.domain.entities.club import Club
from src.domain.entities.insurance import Insurance, InsuranceType
from src.domain.entities.license import License
from src.domain.entities.member import Member
from src.infrastructure.web.dependencies import (
    get_all_insurances_use_case,
    get_all_licenses_use_case,
    get_all_members_use_case,
    get_auth_context,
    get_club_repository,
    get_member_repository,
)
from src.infrastructure.web.routers.import_export import router
from tests.api.callers import OWN_CLUB, super_admin

pytestmark = [pytest.mark.api]

CLUB_NAME = "Dojo Example"
MEMBER = Member(id="member-ana", first_name="Ana", last_name="Example Test", dni="00000000T", club_id=OWN_CLUB)
LICENSE = License(id="license-ana", member_id=MEMBER.id, license_number="LIC-1", grade="1 kyu")
INSURANCE = Insurance(id="insurance-ana", member_id=MEMBER.id, insurance_type=InsuranceType.ACCIDENT,
                      policy_number="POL-1", insurance_company="Aseguradora Example")


@pytest.fixture
def club_repository():
    return SimpleNamespace(find_all=AsyncMock(return_value=[Club(id=OWN_CLUB, name=CLUB_NAME, email="dojo@example.com")]))


@pytest.fixture
def client(club_repository) -> TestClient:
    member_repository = SimpleNamespace(find_by_id=AsyncMock(return_value=MEMBER))
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.dependency_overrides = {
        get_auth_context: lambda: super_admin(),
        get_club_repository: lambda: club_repository,
        get_member_repository: lambda: member_repository,
        get_all_members_use_case: lambda: SimpleNamespace(execute=AsyncMock(return_value=[MEMBER])),
        get_all_licenses_use_case: lambda: SimpleNamespace(execute=AsyncMock(return_value=[LICENSE])),
        get_all_insurances_use_case: lambda: SimpleNamespace(execute=AsyncMock(return_value=[INSURANCE])),
    }
    return TestClient(app)


def club_cell(client: TestClient, path: str):
    sheet = openpyxl.load_workbook(BytesIO(client.get(f"/api/v1/import-export/{path}/export").content)).active
    headers = [cell.value for cell in sheet[1]]
    return sheet.cell(row=2, column=headers.index("Club") + 1).value


def test_members_export_writes_the_club_name(client):
    assert club_cell(client, "members") == CLUB_NAME


def test_licences_export_writes_the_club_name(client):
    assert club_cell(client, "licenses") == CLUB_NAME


def test_insurances_export_writes_the_club_name(client):
    assert club_cell(client, "insurances") == CLUB_NAME


def test_a_club_name_that_looks_like_a_formula_is_exported_as_text(client, club_repository):
    formula = '=HYPERLINK("https://example.com","x")'
    club_repository.find_all.return_value = [Club(id=OWN_CLUB, name=formula, email="dojo@example.com")]

    sheet = openpyxl.load_workbook(BytesIO(client.get("/api/v1/import-export/members/export").content)).active
    headers = [cell.value for cell in sheet[1]]
    cell = sheet.cell(row=2, column=headers.index("Club") + 1)

    assert (cell.value, cell.data_type) == (formula, "s")
